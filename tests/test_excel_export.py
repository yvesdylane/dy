from io import BytesIO

from openpyxl import load_workbook

from utils.excel_export import build_excel

HEADERS = [
    "Name",
    "Surname",
    "Phone",
    "Email",
    "Gender",
    "Department",
    "Group",
    "School",
    "DOB",
    "Quarter",
    "Fees Paid",
    "Total Fees",
    "Active",
    "Created At",
]


def test_build_excel_writes_header_and_rows():
    rows = [
        {
            "name": "Alice",
            "surname": "Ngo",
            "phone": "+237600000001",
            "email": "alice@example.com",
            "gender": "female",
            "department": "SWE",
            "group": "A",
            "school": "UB",
            "dob": "2002-01-15",
            "quarter": "1",
            "fees_paid": 20000,
            "total_fees": 40000,
            "is_active": True,
            "created_at": "2026-01-01T00:00:00",
        }
    ]

    wb = load_workbook(BytesIO(build_excel(rows)))
    ws = wb["Interns"]

    assert [c.value for c in ws[1]] == HEADERS
    assert [c.value for c in ws[2]] == [
        "Alice",
        "Ngo",
        "+237600000001",
        "alice@example.com",
        "female",
        "SWE",
        "A",
        "UB",
        "2002-01-15",
        "1",
        20000,
        40000,
        True,
        "2026-01-01T00:00:00",
    ]


def test_build_excel_with_empty_rows_has_header_only():
    wb = load_workbook(BytesIO(build_excel([])))
    ws = wb["Interns"]

    assert ws.max_row == 1
    assert [c.value for c in ws[1]] == HEADERS
